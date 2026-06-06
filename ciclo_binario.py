from refprop_utils import rprop, WATER_CONFIG, init_refprop
import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
from concurrent.futures import ProcessPoolExecutor
from tqdm import tqdm
import os
import json
from typing import Any

# ---------------------------------------------------------------------------
# Columnas del DataFrame de resultados
# ---------------------------------------------------------------------------
# fluid_A, fluid_B, x_A, x_B, water_config,
# t_1, p_1, h_1, s_1, d_1,
# t_2, p_2, h_2, s_2, d_2,
# t_3, p_3, h_3, s_3, d_3,
# t_4, p_4, h_4, s_4, d_4,
# COP, VHC, pinch, glide_k, glide_0,
# approach_k, error

COLS = [
    "fluid_A", "fluid_B", "x_A", "x_B", "water_config",
    "t_1", "p_1", "h_1", "s_1", "d_1",
    "t_2", "p_2", "h_2", "s_2", "d_2",
    "t_3", "p_3", "h_3", "s_3", "d_3",
    "t_4", "p_4", "h_4", "s_4", "d_4",
    "COP", "VHC", "pinch", "glide_k", "glide_0",
    "approach_k", "error",
]


def _fila_error(fluid_A, fluid_B, x_A, x_B, water_config, error: str) -> dict:
    row = {c: np.nan for c in COLS}
    row.update(
        fluid_A=fluid_A,
        fluid_B=fluid_B,
        x_A=x_A,
        x_B=x_B,
        water_config=water_config,
        error=error,
    )
    return row


def _fila_ok(
    fluid_A, fluid_B, x_A, x_B, water_config, approach_k,
    pts: dict,          # {"1": (t,p,h,s,d), "2": ..., "3": ..., "4": ...}
    COP, VHC, pinch, glide_k, glide_0,
) -> dict:
    row: dict = {}
    row["fluid_A"] = fluid_A
    row["fluid_B"] = fluid_B
    row["x_A"] = x_A
    row["x_B"] = x_B
    row["water_config"] = water_config
    for n, (t, p, h, s, d) in pts.items():
        row[f"t_{n}"] = t
        row[f"p_{n}"] = p
        row[f"h_{n}"] = h
        row[f"s_{n}"] = s
        row[f"d_{n}"] = d
    row["COP"] = COP
    row["VHC"] = VHC
    row["pinch"] = pinch
    row["glide_k"] = glide_k
    row["glide_0"] = glide_0
    row["approach_k"] = approach_k
    row["error"] = None
    return row


# ---------------------------------------------------------------------------
# Ciclo básico: una sola llamada
# ---------------------------------------------------------------------------

def calcular_ciclo(
    fluid_A: str,
    fluid_B: str | None,
    x_A: float,
    water_config: str,
    approach_k: float,
) -> dict:
    """
    Calcula un punto del ciclo. Devuelve un dict con las columnas de COLS.
    Si fluid_B es None, se trata como fluido puro (x_A = 1).
    """
    fluido = [fluid_A, fluid_B] if fluid_B else fluid_A
    mezcla = [x_A, 1 - x_A] if fluid_B else [1.0]
    x_B = round(1 - x_A, 10) if fluid_B else 0.0

    temperaturas_agua = WATER_CONFIG[water_config]
    t_hw_in, t_hw_out = temperaturas_agua["t_hw"]
    t_cw_in, t_cw_out = temperaturas_agua["t_cw"]

    ap_k = approach_k
    ap_0 = 3
    SH = 5
    SUB = 1
    rend_iso_h = 0.6

    try:
        t3 = t_hw_in + ap_k
        T_crit = rprop(fluido, "Tcrit", mezcla, T=0, H=0)
        if t3 > T_crit:
            return _fila_error(fluid_A, fluid_B, x_A, x_B, water_config, "Transcrítico")

        PK = rprop(fluido, "P", mezcla, T=t3 + SUB, Q=0)

        # Punto 3
        [h3, d3] = rprop(fluido, ["H", "D"], mezcla, T=t3, P=PK)
        p3 = PK

        # Punto 4
        [p4, h4, t4] = rprop(fluido, ["P", "H", "T"], mezcla, H=h3, T=t_cw_out - ap_0)
        P0 = p4

        # Punto 1
        t_sat_1 = rprop(fluido, "T", mezcla, P=P0, Q=1)
        t1 = t_sat_1 + SH
        [h1, s1, d1] = rprop(fluido, ["H", "S", "D"], mezcla, T=t1, P=P0)
        p1 = P0

        # Punto 2
        h2_s = rprop(fluido, "H", mezcla, P=PK, S=s1)
        h2 = h1 + (h2_s - h1) / rend_iso_h
        [q2, d2, t2] = rprop(fluido, ["Q", "D", "T"], mezcla, P=PK, H=h2)
        p2 = PK
        if q2 <= 1:
            return _fila_error(fluid_A, fluid_B, x_A, x_B, water_config, "Bifásico")

        # s2 (solo para completar el punto)
        s2 = rprop(fluido, "S", mezcla, P=PK, H=h2)
        s3 = rprop(fluido, "S", mezcla, T=t3, P=PK)
        s4 = rprop(fluido, "S", mezcla, H=h4, P=p4)
        d4 = rprop(fluido, "D", mezcla, H=h4, P=p4)

        # COP y VHC
        COP = (h2 - h3) / (h2 - h1)
        VHC = d1 * (h2 - h3)

        # Puntos saturados en PK
        t_pk_liq = rprop(fluido, "T", mezcla, P=PK, Q=0)
        t_pk_vap = rprop(fluido, "T", mezcla, P=PK, Q=1)
        h_pk_vap = rprop(fluido, "H", mezcla, P=PK, Q=1)

        # Puntos saturados en P0
        t_p0_liq = rprop(fluido, "T", mezcla, P=P0, Q=0)
        t_p0_vap = rprop(fluido, "T", mezcla, P=P0, Q=1)

        # Pinch
        h_water_out = rprop("WATER", "H", [1.0], T=t_hw_out, P=1)
        h_water_in  = rprop("WATER", "H", [1.0], T=t_hw_in,  P=1)
        ratio_m = (h2 - h3) / (h_water_out - h_water_in)
        h_water_pinch = h_water_out - (1 / ratio_m) * (h2 - h_pk_vap)
        t_water_pinch = rprop("WATER", "T", [1.0], H=h_water_pinch, P=1)
        pinch = t_pk_vap - t_water_pinch

        # Glide
        glide_k = t_pk_vap - t_pk_liq
        glide_0 = t_p0_vap - t4   # igual que en el código original (T del punto 4)

        pts = {
            "1": (t1, p1, h1, s1, d1),
            "2": (t2, p2, h2, s2, d2),
            "3": (t3, p3, h3, s3, d3),
            "4": (t4, p4, h4, s4, d4),
        }

        return _fila_ok(
            fluid_A, fluid_B, x_A, x_B, water_config, ap_k,
            pts, COP, VHC, pinch, glide_k, glide_0,
        )

    except ZeroDivisionError:
        return _fila_error(fluid_A, fluid_B, x_A, x_B, water_config, "División 0")
    except RuntimeError:
        return _fila_error(fluid_A, fluid_B, x_A, x_B, water_config, "REFPROP")


# ---------------------------------------------------------------------------
# Ciclo básico con búsqueda de approach
# ---------------------------------------------------------------------------

def calcular_ciclo_basico(
    fluid_A: str,
    fluid_B: str | None,
    x_A: float,
    water_config: str,
    approach_ini: float = 6.5,
    approach_max: float = 20,
    step: float = 0.5,
) -> dict:
    approach = approach_ini
    last_row = None

    while approach < approach_max:
        row = calcular_ciclo(fluid_A, fluid_B, x_A, water_config, approach)
        last_row = row

        if row["error"] is not None:
            return row
        if row["pinch"] >= 1:
            return row

        approach += step

    # Nunca se cumplió pinch >= 1
    x_B = 1 - x_A
    row = last_row or _fila_error(fluid_A, fluid_B, x_A, x_B if fluid_B else 0.0, water_config, "PinchBajo")
    row["error"] = "PinchBajo"
    return row


# ---------------------------------------------------------------------------
# Valores de referencia (propano)
# ---------------------------------------------------------------------------

def calcular_valores_referencia(water_config: str) -> tuple[float, float, float]:
    margen = 0.3
    row_propano = calcular_ciclo_basico("PROPANE", None, 1.0, water_config)
    vhc_ref = row_propano["VHC"]
    cop_ref = row_propano["COP"]
    return (1 - margen) * vhc_ref, (1 + margen) * vhc_ref, cop_ref


# ---------------------------------------------------------------------------
# Filtrado sobre DataFrame
# ---------------------------------------------------------------------------

def filtrar(df: pd.DataFrame, vhc_min: float, vhc_max: float) -> pd.DataFrame:
    mask = (
        df["error"].isna()
        & df["VHC"].between(vhc_min, vhc_max)
        & (df["t_2"] < 130)
        & (df["p_2"] < 28)
        & (df["glide_k"] < 10)
        & (df["glide_0"] < 10)
    )
    return df[mask].sort_values("COP", ascending=False)


# ---------------------------------------------------------------------------
# Worker para ProcessPoolExecutor
# ---------------------------------------------------------------------------

def _worker(args):
    import refprop_utils
    if refprop_utils.RP is None:
        raise RuntimeError("REFPROP no inicializado en el worker")
    fluid_A, fluid_B, x_A, water_config = args
    return calcular_ciclo_basico(fluid_A, fluid_B, x_A, water_config)


# ---------------------------------------------------------------------------
# Cálculo bruto de mezclas binarias
# ---------------------------------------------------------------------------

def calcular_mezclas(
    posibles_refrigerantes: list[str],
    water_config: str,
) -> pd.DataFrame:
    n_calcs = 41
    props_a = [float(x) for x in np.linspace(0, 1, n_calcs)]

    n = len(posibles_refrigerantes)
    total = n * (n - 1) // 2

    filas: list[dict] = []

    with tqdm(total=total, desc="Calculando mezclas") as pbar:
        for i, ref_a in enumerate(posibles_refrigerantes[:-1]):
            for ref_b in posibles_refrigerantes[i + 1:]:

                inputs = [
                    (ref_a, ref_b, x_a, water_config)
                    for x_a in props_a
                ]

                cpu = os.cpu_count() // 2 or 1
                with ProcessPoolExecutor(max_workers=cpu, initializer=init_refprop) as ex:
                    resultados = list(ex.map(_worker, inputs, chunksize=2))

                filas.extend(resultados)
                pbar.update(1)

    df = pd.DataFrame(filas, columns=COLS)

    # Guardar como parquet (más eficiente que JSON para DataFrames)
    out_dir = os.path.join("resultados_ciclo_basico", water_config, "binarias")
    os.makedirs(out_dir, exist_ok=True)
    df.to_parquet(os.path.join(out_dir, "resultados.parquet"), index=False)

    return df


# ---------------------------------------------------------------------------
# Refinado fino
# ---------------------------------------------------------------------------

def refinar_mezclas(water_config: str) -> pd.DataFrame:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados.parquet")
    df = pd.read_parquet(path_parquet)

    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)

    salto_fino = 0.005
    salto_busq = 0.025

    filas_finas: list[dict] = []

    pares_vistos: set[tuple] = set()

    for (ref_a, ref_b), grupo in df.groupby(["fluid_A", "fluid_B"]):
        par = tuple(sorted([ref_a, ref_b]))
        if par in pares_vistos:
            continue
        pares_vistos.add(par)

        candidatos = filtrar(grupo, vhc_min, vhc_max).head(2)
        if candidatos.empty:
            continue

        # Determinar rangos de composición a explorar
        comps_rangos: list[tuple[float, float]] = []

        if len(candidatos) == 2:
            x0, x1 = candidatos.iloc[0]["x_A"], candidatos.iloc[1]["x_A"]
            if abs(x0 - x1) <= 0.10:
                comps_rangos.append((min(x0, x1), max(x0, x1)))
            else:
                comps_rangos.append((max(x0 - salto_busq, 0), min(x0 + salto_busq, 1)))
                comps_rangos.append((max(x1 - salto_busq, 0), min(x1 + salto_busq, 1)))
        else:
            x0 = candidatos.iloc[0]["x_A"]
            comps_rangos.append((max(x0 - salto_busq, 0), min(x0 + salto_busq, 1)))

        for (lo, hi) in comps_rangos:
            n_steps = max(int(round((hi - lo) / salto_fino)), 1) + 1
            for x_a in np.linspace(lo, hi, n_steps):
                row = calcular_ciclo_basico(ref_a, ref_b, float(x_a), water_config)
                filas_finas.append(row)

    df_fino = pd.DataFrame(filas_finas, columns=COLS)

    out_dir = os.path.join("resultados_ciclo_basico", water_config, "binarias")
    df_fino.to_parquet(os.path.join(out_dir, "resultados_finos.parquet"), index=False)

    return df_fino


# ---------------------------------------------------------------------------
# Excel: resultados brutos
# ---------------------------------------------------------------------------

def df_a_excel(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados.xlsx")

    df = pd.read_parquet(path_parquet)

    cols_mostrar = [
    "fluid_A", "fluid_B", "x_A", "x_B", "water_config",
    "t_1", "p_1", "h_1", "s_1", "d_1",
    "t_2", "p_2", "h_2", "s_2", "d_2",
    "t_3", "p_3", "h_3", "s_3", "d_3",
    "t_4", "p_4", "h_4", "s_4", "d_4",
    "COP", "VHC", "pinch", "glide_k", "glide_0",
    "approach_k", "error",
    ]

    df[cols_mostrar].to_excel(path_excel, index=False)

    wb = load_workbook(path_excel)
    ws = wb.active
    align = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = "A2"
    wb.save(path_excel)


# ---------------------------------------------------------------------------
# Excel: resultados filtrados
# ---------------------------------------------------------------------------

def df_a_excel_filtrado(water_config: str) -> None:
    path_parquet  = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados.parquet")
    path_excel    = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados_filtrados.xlsx")

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, _ = calcular_valores_referencia(water_config)
    df_f = filtrar(df, vhc_min, vhc_max)

    cols_mostrar = [
    "fluid_A", "fluid_B", "x_A", "x_B", "water_config",
    "t_1", "p_1", "h_1", "s_1", "d_1",
    "t_2", "p_2", "h_2", "s_2", "d_2",
    "t_3", "p_3", "h_3", "s_3", "d_3",
    "t_4", "p_4", "h_4", "s_4", "d_4",
    "COP", "VHC", "pinch", "glide_k", "glide_0",
    "approach_k", "error",
    ]

    df_f[cols_mostrar].to_excel(path_excel, index=False)

    wb = load_workbook(path_excel)
    ws = wb.active
    align = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = "A2"
    wb.save(path_excel)


# ---------------------------------------------------------------------------
# Excel: resultados finos
# ---------------------------------------------------------------------------

def df_a_excel_fino(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados_finos.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados_finos.xlsx")

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, _ = calcular_valores_referencia(water_config)
    df_f = filtrar(df, vhc_min, vhc_max)

    cols_mostrar = [
    "fluid_A", "fluid_B", "x_A", "x_B", "water_config",
    "COP", "VHC", "pinch", "glide_k", "glide_0",
    "approach_k", "error",
    "t_1", "p_1", "h_1", "s_1", "d_1",
    "t_2", "p_2", "h_2", "s_2", "d_2",
    "t_3", "p_3", "h_3", "s_3", "d_3",
    "t_4", "p_4", "h_4", "s_4", "d_4"
    ]

    df_f[cols_mostrar].to_excel(path_excel, index=False)

    wb = load_workbook(path_excel)
    ws = wb.active
    align = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = "A2"
    wb.save(path_excel)


# ---------------------------------------------------------------------------
# Excel resumen (mejor COP por par, solo mezclas con COP >= propano)
# ---------------------------------------------------------------------------

def crear_excel_resumen(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados_finos.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resumen_resultados.xlsx")
 
    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)
    vhc_propano = (vhc_min + vhc_max) / 2
    df_f = filtrar(df, vhc_min, vhc_max)
    df_f = df_f[df_f["COP"] >= cop_propano]
 
    cols_mostrar = [
        "fluid_A", "fluid_B", "x_A", "x_B", "water_config",
        "COP", "ΔCOP", "VHC", "ΔVHC", "pinch", "glide_k", "glide_0",
        "approach_k", "error",
        "t_1", "p_1", "h_1", "s_1", "d_1",
        "t_2", "p_2", "h_2", "s_2", "d_2",
        "t_3", "p_3", "h_3", "s_3", "d_3",
        "t_4", "p_4", "h_4", "s_4", "d_4",
    ]
 
    def add_deltas(dataframe):
        dataframe = dataframe.copy()
        dataframe["ΔCOP"] = (dataframe["COP"] - cop_propano)/cop_propano
        dataframe["ΔVHC"] = (dataframe["VHC"] - vhc_propano)/vhc_propano
        return dataframe
 
    def best_per_pair(dataframe):
        return (
            dataframe.sort_values("COP", ascending=False)
                     .groupby(["fluid_A", "fluid_B"], sort=False)
                     .first()
                     .reset_index()
        )
 
    def select_cols(dataframe):
        cols = [c for c in cols_mostrar if c in dataframe.columns]
        return dataframe[cols]
 
    df_sin_filtro = select_cols(add_deltas(best_per_pair(df)))
    df_con_filtro = select_cols(add_deltas(best_per_pair(df_f))) if not df_f.empty else pd.DataFrame(columns=cols_mostrar)
 
    os.makedirs(os.path.dirname(path_excel), exist_ok=True)
 
    with pd.ExcelWriter(path_excel, engine="openpyxl") as writer:
        df_sin_filtro.to_excel(writer, sheet_name="sin_filtrar", index=False)
        df_con_filtro.to_excel(writer, sheet_name="filtrado",    index=False)
 
    wb = load_workbook(path_excel)
    align = Alignment(horizontal="center", vertical="center")
    bold  = Font(bold=True)
 
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.alignment = align
            cell.font = bold
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 15
            for cell in col:
                cell.alignment = align
        ws.freeze_panes = "A2"
 
    wb.save(path_excel)


# ---------------------------------------------------------------------------
# Gráficos binarios
# ---------------------------------------------------------------------------

def generar_graficos_binarios(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "binarias", "resultados.parquet")
    output_folder = os.path.join("resultados_ciclo_basico", water_config, "binarias", "graficos")
    os.makedirs(output_folder, exist_ok=True)

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)
    # df_f = filtrar(df, vhc_min, vhc_max)
    df_f = df # Temporal porque no quiero filtrar de momento por VHC

    for (ref_a, ref_b), grupo in df_f.groupby(["fluid_A", "fluid_B"]):
        nombre = f"{ref_a}, {ref_b}"
        grupo = grupo.sort_values("x_A")

        x_values = grupo["x_A"].tolist()
        y_percent = [((cop - cop_propano) / cop_propano) * 100 for cop in grupo["COP"]]

        plt.figure(figsize=(10, 6))
        plt.scatter(x_values, y_percent, color="b", label=f"Desviación COP")
        plt.axhline(y=0, color="r", linestyle="--", label=f"Propano ({cop_propano:.3f})")
        plt.xticks([i / 10 for i in range(11)], [f"{i * 10}%" for i in range(11)])
        plt.gca().yaxis.set_major_formatter(mtick.PercentFormatter(decimals=0))
        plt.title(f"Desviación COP respecto a propano: {nombre}", fontsize=12)
        plt.xlabel(f"Composición de {ref_a}")
        plt.ylabel("Diferencia de COP respecto a propano (%)")
        plt.grid(True, linestyle=":", alpha=0.6)
        plt.legend()
        plt.xlim(0, 1)
        plt.savefig(os.path.join(output_folder, f"{nombre}.png"))
        plt.close()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    init_refprop()

    water_config = "alta"  # "media_7" / "alta_7" / "media" / "alta"
    posibles_refrigerantes = ["PROPANE", "DME", "ISOBUTANE", "PROPYLENE", "CO2", "BUTANE", "ETHANE", "ETHYLENE", "METHANE"]

    # 1. Cálculo bruto
    calcular_mezclas(posibles_refrigerantes, water_config)
    df_a_excel(water_config)

    # 2. Refinado fino
    refinar_mezclas(water_config)
    df_a_excel_fino(water_config)

    # 3. Gráficos y filtrado
    df_a_excel_filtrado(water_config)
    generar_graficos_binarios(water_config)

    # 4. Resumen
    crear_excel_resumen(water_config)


if __name__ == "__main__":
    main()
