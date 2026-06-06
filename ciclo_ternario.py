import matplotlib.pyplot as plt
import matplotlib.colors as mcolors
import ternary
import numpy as np
import os
from concurrent.futures import ProcessPoolExecutor
from tqdm import tqdm
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from scipy.interpolate import griddata

from refprop_utils import rprop, WATER_CONFIG, init_refprop

# ---------------------------------------------------------------------------
# Columnas del DataFrame de resultados ternarios
# ---------------------------------------------------------------------------

COLS_T = [
    "fluid_A", "fluid_B", "fluid_C", "x_A", "x_B", "x_C", "water_config",
    "t_1", "p_1", "h_1", "s_1", "d_1",
    "t_2", "p_2", "h_2", "s_2", "d_2",
    "t_3", "p_3", "h_3", "s_3", "d_3",
    "t_4", "p_4", "h_4", "s_4", "d_4",
    "COP", "VHC", "pinch", "glide_k", "glide_0",
    "approach_k", "error",
]


# ---------------------------------------------------------------------------
# Helpers internos de fila
# ---------------------------------------------------------------------------

def _fila_error(
    fluid_A: str, fluid_B: str, fluid_C: str,
    x_A: float, x_B: float, x_C: float,
    water_config: str, error: str,
) -> dict:
    row = {c: np.nan for c in COLS_T}
    row.update(
        fluid_A=fluid_A, fluid_B=fluid_B, fluid_C=fluid_C,
        x_A=x_A, x_B=x_B, x_C=x_C,
        water_config=water_config, error=error,
    )
    return row


def _fila_ok(
    fluid_A: str, fluid_B: str, fluid_C: str,
    x_A: float, x_B: float, x_C: float,
    water_config: str, approach_k: float,
    pts: dict,
    COP: float, VHC: float, pinch: float, glide_k: float, glide_0: float,
) -> dict:
    row: dict = {}
    row["fluid_A"]      = fluid_A
    row["fluid_B"]      = fluid_B
    row["fluid_C"]      = fluid_C
    row["x_A"]          = x_A
    row["x_B"]          = x_B
    row["x_C"]          = x_C
    row["water_config"] = water_config
    for n, (t, p, h, s, d) in pts.items():
        row[f"t_{n}"] = t
        row[f"p_{n}"] = p
        row[f"h_{n}"] = h
        row[f"s_{n}"] = s
        row[f"d_{n}"] = d
    row["COP"]       = COP
    row["VHC"]       = VHC
    row["pinch"]     = pinch
    row["glide_k"]   = glide_k
    row["glide_0"]   = glide_0
    row["approach_k"]= approach_k
    row["error"]     = None
    return row


# ---------------------------------------------------------------------------
# Ciclo básico: una sola llamada (3 componentes)
# ---------------------------------------------------------------------------

def calcular_ciclo(
    fluid_A: str,
    fluid_B: str,
    fluid_C: str,
    x_A: float,
    x_B: float,
    water_config: str,
    approach_k: float,
) -> dict:
    """
    Calcula un punto del ciclo ternario.
    Devuelve un dict con las columnas de COLS_T.
    x_C se deriva como 1 - x_A - x_B.
    """
    x_C = round(1 - x_A - x_B, 10)

    if not (0 <= x_C <= 1):
        return _fila_error(fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
                           water_config, "ComposiciónInválida")

    fluido = [fluid_A, fluid_B, fluid_C]
    mezcla = [x_A, x_B, x_C]

    temperaturas_agua = WATER_CONFIG[water_config]
    t_hw_in, t_hw_out = temperaturas_agua["t_hw"]
    t_cw_in, t_cw_out = temperaturas_agua["t_cw"]

    ap_k = approach_k
    ap_0 = 3
    SH = 5
    SUB = 1
    rend_iso_h = 0.6

    try:
        t3 = t_hw_in + ap_k
        T_crit = rprop(fluido, "Tcrit", mezcla, T=0, H=0)
        if t3 > T_crit:
            return _fila_error(fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
                               water_config, "Transcrítico")

        PK = rprop(fluido, "P", mezcla, T=t3 + SUB, Q=0)

        # Punto 3
        [h3, d3] = rprop(fluido, ["H", "D"], mezcla, T=t3, P=PK)
        p3 = PK

        # Punto 4
        [p4, h4, t4] = rprop(fluido, ["P", "H", "T"], mezcla, H=h3, T=t_cw_out - ap_0)
        P0 = p4

        # Punto 1
        t_sat_1 = rprop(fluido, "T", mezcla, P=P0, Q=1)
        t1 = t_sat_1 + SH
        [h1, s1, d1] = rprop(fluido, ["H", "S", "D"], mezcla, T=t1, P=P0)
        p1 = P0

        # Punto 2
        h2_s = rprop(fluido, "H", mezcla, P=PK, S=s1)
        h2 = h1 + (h2_s - h1) / rend_iso_h
        [q2, d2, t2] = rprop(fluido, ["Q", "D", "T"], mezcla, P=PK, H=h2)
        p2 = PK
        if q2 <= 1:
            return _fila_error(fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
                               water_config, "Bifásico")

        s2 = rprop(fluido, "S", mezcla, P=PK, H=h2)
        s3 = rprop(fluido, "S", mezcla, T=t3,  P=PK)
        s4 = rprop(fluido, "S", mezcla, H=h4,  P=p4)
        d4 = rprop(fluido, "D", mezcla, H=h4,  P=p4)

        # COP y VHC
        COP = (h2 - h3) / (h2 - h1)
        VHC = d1 * (h2 - h3)

        # Puntos saturados en PK
        t_pk_liq = rprop(fluido, "T", mezcla, P=PK, Q=0)
        t_pk_vap = rprop(fluido, "T", mezcla, P=PK, Q=1)
        h_pk_vap = rprop(fluido, "H", mezcla, P=PK, Q=1)

        # Puntos saturados en P0
        t_p0_liq = rprop(fluido, "T", mezcla, P=P0, Q=0)
        t_p0_vap = rprop(fluido, "T", mezcla, P=P0, Q=1)

        # Pinch
        h_water_out   = rprop("WATER", "H", [1.0], T=t_hw_out, P=1)
        h_water_in    = rprop("WATER", "H", [1.0], T=t_hw_in,  P=1)
        ratio_m       = (h2 - h3) / (h_water_out - h_water_in)
        h_water_pinch = h_water_out - (1 / ratio_m) * (h2 - h_pk_vap)
        t_water_pinch = rprop("WATER", "T", [1.0], H=h_water_pinch, P=1)
        pinch         = t_pk_vap - t_water_pinch

        # Glide
        glide_k = t_pk_vap - t_pk_liq
        glide_0 = t_p0_vap - t4

        pts = {
            "1": (t1, p1, h1, s1, d1),
            "2": (t2, p2, h2, s2, d2),
            "3": (t3, p3, h3, s3, d3),
            "4": (t4, p4, h4, s4, d4),
        }

        return _fila_ok(
            fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
            water_config, ap_k,
            pts, COP, VHC, pinch, glide_k, glide_0,
        )

    except ZeroDivisionError:
        return _fila_error(fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
                           water_config, "División 0")
    except RuntimeError:
        return _fila_error(fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
                           water_config, "REFPROP")


# ---------------------------------------------------------------------------
# Ciclo básico con búsqueda de approach (3 componentes)
# ---------------------------------------------------------------------------

def calcular_ciclo_basico(
    fluid_A: str,
    fluid_B: str,
    fluid_C: str,
    x_A: float,
    x_B: float,
    water_config: str,
    approach_ini: float = 6.5,
    approach_max: float = 20,
    step: float = 0.5,
) -> dict:
    approach = approach_ini
    last_row = None

    while approach < approach_max:
        row = calcular_ciclo(fluid_A, fluid_B, fluid_C, x_A, x_B,
                             water_config, approach)
        last_row = row

        if row["error"] is not None:
            return row
        if row["pinch"] >= 1:
            return row

        approach += step

    # Nunca se cumplió pinch >= 1
    x_C = round(1 - x_A - x_B, 10)
    row = last_row or _fila_error(fluid_A, fluid_B, fluid_C, x_A, x_B, x_C,
                                  water_config, "PinchBajo")
    row["error"] = "PinchBajo"
    return row


# ---------------------------------------------------------------------------
# Valores de referencia (propano puro)
# ---------------------------------------------------------------------------

def calcular_valores_referencia(water_config: str) -> tuple[float, float, float]:
    """Calcula VHC y COP de referencia usando propano puro."""
    margen = 0.3
    # Propano puro: x_A=1, x_B=0 → x_C=0. Los tres fluidos se pasan como
    # PROPANE pero REFPROP solo usará la fracción x_A=1.
    row_propano = calcular_ciclo_basico(
        "PROPANE", "PROPANE", "PROPANE", 1.0, 0.0, water_config
    )
    vhc_ref = row_propano["VHC"]
    cop_ref = row_propano["COP"]
    return (1 - margen) * vhc_ref, (1 + margen) * vhc_ref, cop_ref


# ---------------------------------------------------------------------------
# Filtrado sobre DataFrame
# ---------------------------------------------------------------------------

def filtrar_ternario(df: pd.DataFrame, vhc_min: float, vhc_max: float) -> pd.DataFrame:
    mask = (
        df["error"].isna()
        & df["VHC"].between(vhc_min, vhc_max)
        & (df["t_2"] < 130)
        & (df["p_2"] < 28)
        & (df["glide_k"] < 10)
        & (df["glide_0"] < 10)
    )
    return df[mask].sort_values("COP", ascending=False)


# ---------------------------------------------------------------------------
# Worker para ProcessPoolExecutor
# ---------------------------------------------------------------------------

def _worker_ternario(args):
    import refprop_utils
    if refprop_utils.RP is None:
        raise RuntimeError("REFPROP no inicializado en el worker")
    fluid_A, fluid_B, fluid_C, x_A, x_B, water_config = args
    return calcular_ciclo_basico(fluid_A, fluid_B, fluid_C, x_A, x_B, water_config)


# ---------------------------------------------------------------------------
# Helpers de combinaciones
# ---------------------------------------------------------------------------

def crear_lista_3_ref(posibles_refrigerantes: list[str]) -> list[tuple[str, str, str]]:
    lista: list[tuple[str, str, str]] = []
    refs = posibles_refrigerantes
    for i, ref_a in enumerate(refs[:-2]):
        for j, ref_b in enumerate(refs[i + 1:-1]):
            for ref_c in refs[i + j + 2:]:
                lista.append((ref_a, ref_b, ref_c))
    return lista


def crear_props_3_ref(n_prop: int) -> list[tuple[float, float, float]]:
    """Genera todas las combinaciones (x_A, x_B, x_C) con x_A+x_B+x_C=1."""
    props = [float(x) for x in np.linspace(0, 1, n_prop)]
    lista: list[tuple[float, float, float]] = []
    for i, prop_a in enumerate(props):
        limite_b = props if i == 0 else props[:-i]
        for prop_b in limite_b:
            prop_c = round(1 - prop_a - prop_b, 10)
            if prop_c < 0:
                continue
            lista.append((prop_a, prop_b, prop_c))
    return lista


# ---------------------------------------------------------------------------
# Cálculo bruto
# ---------------------------------------------------------------------------

def calcular_resultados(
    posibles_refrigerantes: list[str],
    water_config: str,
    n_prop: int,
) -> pd.DataFrame:

    combinaciones_ref = crear_lista_3_ref(posibles_refrigerantes)
    rango_proporciones = crear_props_3_ref(n_prop)

    lista_inputs = [
        (ref_a, ref_b, ref_c, x_A, x_B, water_config)
        for (ref_a, ref_b, ref_c) in combinaciones_ref
        for (x_A, x_B, _) in rango_proporciones
    ]

    cpu = os.cpu_count() // 2 or 1
    chunksize = 2

    with ProcessPoolExecutor(max_workers=cpu, initializer=init_refprop) as ex:
        filas = list(tqdm(
            ex.map(_worker_ternario, lista_inputs, chunksize=chunksize),
            total=len(lista_inputs),
            desc="Calculando mezclas ternarias (cálculo bruto)",
        ))

    df = pd.DataFrame(filas, columns=COLS_T)

    out_dir = os.path.join("resultados_ciclo_basico", water_config, "ternarias")
    os.makedirs(out_dir, exist_ok=True)
    df.to_parquet(os.path.join(out_dir, "resultados.parquet"), index=False)

    return df


# ---------------------------------------------------------------------------
# Rango de composiciones fino (equivalente a crear_rango_composiciones)
# Devuelve lista de (x_A, x_B) a recalcular alrededor de los candidatos
# ---------------------------------------------------------------------------

def _crear_rango_composiciones_ternario(
    candidatos: pd.DataFrame,
) -> list[tuple[float, float]]:
    """
    A partir de los candidatos (máx 2 filas del DF filtrado),
    devuelve una lista plana de (x_A, x_B) para el cálculo fino.
    """
    props = list(zip(candidatos["x_A"].tolist(), candidatos["x_B"].tolist()))
    dx = dy = (-0.05, 0.05)
    salto = 0.01

    def cuadricula(x_min, x_max, y_min, y_max) -> list[tuple[float, float]]:
        n_x = max(int(round((x_max - x_min) / salto)), 1) + 1
        n_y = max(int(round((y_max - y_min) / salto)), 1) + 1
        puntos = []
        for i in range(n_x):
            for j in range(n_y):
                x = round(x_min + i * salto, 3)
                y = round(y_min + j * salto, 3)
                z = round(1 - x - y, 3)
                if 0 <= x <= 1 and 0 <= y <= 1 and 0 <= z <= 1:
                    puntos.append((x, y))
        return puntos

    todos: list[tuple[float, float]] = []

    if len(props) == 2:
        (x0, y0), (x1, y1) = props
        cerca_x = abs(x0 - x1) <= 0.1
        cerca_y = abs(y0 - y1) <= 0.1

        if not (cerca_x and cerca_y):
            # Lejos: cuadrado separado por cada punto
            for x, y in props:
                todos += cuadricula(
                    max(x + min(dx), 0), min(x + max(dx), 1),
                    max(y + min(dy), 0), min(y + max(dy), 1),
                )
        elif x0 != x1 and y0 != y1:
            # Cerca pero distintos: cuadrado grande
            todos += cuadricula(min(x0, x1), max(x0, x1), min(y0, y1), max(y0, y1))
        elif x0 == x1:
            # Comparten X: ensanchar X, recorrer Y
            todos += cuadricula(
                max(x0 + min(dx), 0), min(x0 + max(dx), 1),
                min(y0, y1), max(y0, y1),
            )
        else:
            # Comparten Y: ensanchar Y, recorrer X
            todos += cuadricula(
                min(x0, x1), max(x0, x1),
                max(y0 + min(dy), 0), min(y0 + max(dy), 1),
            )

    elif len(props) == 1:
        x, y = props[0]
        todos += cuadricula(
            max(x + min(dx), 0), min(x + max(dx), 1),
            max(y + min(dy), 0), min(y + max(dy), 1),
        )

    # Eliminar duplicados preservando orden
    vistos: set[tuple[float, float]] = set()
    unicos: list[tuple[float, float]] = []
    for p in todos:
        if p not in vistos:
            vistos.add(p)
            unicos.append(p)
    return unicos


# ---------------------------------------------------------------------------
# Refinado fino
# ---------------------------------------------------------------------------

def refinar_mezclas(water_config: str) -> pd.DataFrame:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados.parquet")
    df = pd.read_parquet(path_parquet)

    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)

    filas_finas: list[dict] = []
    ternas_vistas: set[tuple[str, str, str]] = set()

    grupos = df.groupby(["fluid_A", "fluid_B", "fluid_C"])
    for (ref_a, ref_b, ref_c), grupo in tqdm(grupos, desc="Refinando mezclas ternarias"):
        terna = tuple(sorted([ref_a, ref_b, ref_c]))
        if terna in ternas_vistas:
            continue
        ternas_vistas.add(terna)

        candidatos = filtrar_ternario(grupo, vhc_min, vhc_max).head(2)
        if candidatos.empty:
            continue

        puntos_finos = _crear_rango_composiciones_ternario(candidatos)
        if not puntos_finos:
            continue

        lista_inputs = [
            (ref_a, ref_b, ref_c, float(x_A), float(x_B), water_config)
            for (x_A, x_B) in puntos_finos
        ]

        cpu = os.cpu_count() // 2 or 1
        with ProcessPoolExecutor(max_workers=cpu, initializer=init_refprop) as ex:
            resultados = list(ex.map(_worker_ternario, lista_inputs, chunksize=2))

        filas_finas.extend(resultados)

    df_fino = pd.DataFrame(filas_finas, columns=COLS_T)

    out_dir = os.path.join("resultados_ciclo_basico", water_config, "ternarias")
    df_fino.to_parquet(os.path.join(out_dir, "resultados_finos.parquet"), index=False)

    return df_fino


# ---------------------------------------------------------------------------
# Excel: resultados brutos
# ---------------------------------------------------------------------------

def df_a_excel(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados.xlsx")

    df = pd.read_parquet(path_parquet)

    cols_mostrar = [
        "fluid_A", "fluid_B", "fluid_C", "x_A", "x_B", "x_C", "water_config",
        "t_1", "p_1", "h_1", "s_1", "d_1",
        "t_2", "p_2", "h_2", "s_2", "d_2",
        "t_3", "p_3", "h_3", "s_3", "d_3",
        "t_4", "p_4", "h_4", "s_4", "d_4",
        "COP", "VHC", "pinch", "glide_k", "glide_0",
        "approach_k", "error",
    ]

    df[cols_mostrar].to_excel(path_excel, index=False)
    _formato_excel(path_excel)


# ---------------------------------------------------------------------------
# Excel: resultados filtrados
# ---------------------------------------------------------------------------

def df_a_excel_filtrado(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados_filtrados.xlsx")

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, _ = calcular_valores_referencia(water_config)
    df_f = filtrar_ternario(df, vhc_min, vhc_max)

    cols_mostrar = [
        "fluid_A", "fluid_B", "fluid_C", "x_A", "x_B", "x_C", "water_config",
        "t_1", "p_1", "h_1", "s_1", "d_1",
        "t_2", "p_2", "h_2", "s_2", "d_2",
        "t_3", "p_3", "h_3", "s_3", "d_3",
        "t_4", "p_4", "h_4", "s_4", "d_4",
        "COP", "VHC", "pinch", "glide_k", "glide_0",
        "approach_k", "error",
    ]

    df_f[cols_mostrar].to_excel(path_excel, index=False)
    _formato_excel(path_excel)


# ---------------------------------------------------------------------------
# Excel: resultados finos
# ---------------------------------------------------------------------------

def df_a_excel_fino(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados_finos.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados_finos.xlsx")

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, _ = calcular_valores_referencia(water_config)
    df_f = filtrar_ternario(df, vhc_min, vhc_max)

    cols_mostrar = [
        "fluid_A", "fluid_B", "fluid_C", "x_A", "x_B", "x_C", "water_config",
        "COP", "VHC", "pinch", "glide_k", "glide_0",
        "approach_k", "error",
        "t_1", "p_1", "h_1", "s_1", "d_1",
        "t_2", "p_2", "h_2", "s_2", "d_2",
        "t_3", "p_3", "h_3", "s_3", "d_3",
        "t_4", "p_4", "h_4", "s_4", "d_4",
    ]

    df_f[cols_mostrar].to_excel(path_excel, index=False)
    _formato_excel(path_excel)


# ---------------------------------------------------------------------------
# Excel resumen (mejor COP por terna, con y sin filtro de COP >= propano)
# ---------------------------------------------------------------------------

def crear_excel_resumen(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados_finos.parquet")
    path_excel   = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resumen_resultados.xlsx")

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)
    vhc_propano = (vhc_min + vhc_max) / 2

    df_f = filtrar_ternario(df, vhc_min, vhc_max)
    df_filtrado_cop = df_f[df_f["COP"] >= cop_propano]

    cols_mostrar = [
        "fluid_A", "fluid_B", "fluid_C", "x_A", "x_B", "x_C", "water_config",
        "COP", "ΔCOP", "VHC", "ΔVHC", "pinch", "glide_k", "glide_0",
        "approach_k", "error",
        "t_1", "p_1", "h_1", "s_1", "d_1",
        "t_2", "p_2", "h_2", "s_2", "d_2",
        "t_3", "p_3", "h_3", "s_3", "d_3",
        "t_4", "p_4", "h_4", "s_4", "d_4",
    ]

    def add_deltas(dataframe: pd.DataFrame) -> pd.DataFrame:
        dataframe = dataframe.copy()
        dataframe["ΔCOP"] = (dataframe["COP"] - cop_propano) / cop_propano
        dataframe["ΔVHC"] = (dataframe["VHC"] - vhc_propano) / vhc_propano
        return dataframe

    def best_per_terna(dataframe: pd.DataFrame) -> pd.DataFrame:
        return (
            dataframe.sort_values("COP", ascending=False)
                     .groupby(["fluid_A", "fluid_B", "fluid_C"], sort=False)
                     .first()
                     .reset_index()
        )

    def select_cols(dataframe: pd.DataFrame) -> pd.DataFrame:
        cols = [c for c in cols_mostrar if c in dataframe.columns]
        return dataframe[cols]

    df_sin_filtro  = select_cols(add_deltas(best_per_terna(df)))
    df_con_filtro  = select_cols(add_deltas(best_per_terna(df_filtrado_cop))) \
        if not df_filtrado_cop.empty else pd.DataFrame(columns=cols_mostrar)

    os.makedirs(os.path.dirname(path_excel), exist_ok=True)

    with pd.ExcelWriter(path_excel, engine="openpyxl") as writer:
        df_sin_filtro.to_excel(writer, sheet_name="sin_filtrar", index=False)
        df_con_filtro.to_excel(writer, sheet_name="filtrado",    index=False)

    wb = load_workbook(path_excel)
    align = Alignment(horizontal="center", vertical="center")
    bold  = Font(bold=True)

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.alignment = align
            cell.font = bold
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 15
            for cell in col:
                cell.alignment = align
        ws.freeze_panes = "A2"

    wb.save(path_excel)


# ---------------------------------------------------------------------------
# TXT resumen (equivalente a guardar_txt del original)
# ---------------------------------------------------------------------------

def guardar_txt(water_config: str) -> None:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados_finos.parquet")
    path_txt     = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados_finos.txt")

    df = pd.read_parquet(path_parquet)
    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)
    df_f = filtrar_ternario(df, vhc_min, vhc_max).sort_values("COP", ascending=False)

    lines: list[str] = []
    for _, row in df_f.iterrows():
        comp = (
            f"{row['fluid_A']}: {row['x_A']*100:.0f}%, "
            f"{row['fluid_B']}: {row['x_B']*100:.0f}%, "
            f"{row['fluid_C']}: {row['x_C']*100:.0f}%"
        )
        proporcion = (row["COP"] / cop_propano - 1) * 100
        if proporcion >= 0:
            lines.append(comp + f"\nCOP {proporcion:.2f}% más GRANDE que el propano\n\n")
        else:
            lines.append(comp + f"\nCOP {-proporcion:.2f}% más PEQUEÑO que el propano\n\n")

    with open(path_txt, "w", encoding="utf-8") as f:
        f.writelines(lines)


# ---------------------------------------------------------------------------
# Gráficos ternarios
# ---------------------------------------------------------------------------

def obtener_casos(water_config: str) -> tuple[list[dict], dict[str, dict[str, Any]]]:
    path_parquet = os.path.join("resultados_ciclo_basico", water_config, "ternarias", "resultados.parquet")
    df = pd.read_parquet(path_parquet)

    vhc_min, vhc_max, cop_propano = calcular_valores_referencia(water_config)
    df_f = filtrar_ternario(df, vhc_min, vhc_max)

    casos: list[dict] = []

    for (ref_a, ref_b, ref_c), grupo in df_f.groupby(["fluid_A", "fluid_B", "fluid_C"]):
        valores: dict[tuple, dict[str, float]] = {}
        for _, row in grupo.iterrows():
            coord = (row["x_A"], row["x_B"], row["x_C"])
            valores[coord] = {
                "COP":        row["COP"],
                "VHC":        row["VHC"],
                "T pinch":    row["pinch"],
                "T descarga": row["t_2"],
                "Presion k":  row["p_2"],
                "Presion 0":  row["p_1"],
                "glide k":    row["glide_k"],
                "glide 0":    row["glide_0"],
            }
        if valores:
            casos.append({
                "nombre": [ref_a, ref_b, ref_c],
                "valores": valores,
            })

    config_mag: dict[str, dict[str, Any]] = {
        "COP":        {"perc": True,  "ref": {"type": "set_center",  "val": cop_propano}},
        "VHC":        {"perc": True,  "ref": {"type": "set_min_max", "val": [vhc_min, vhc_max]}},
        "T pinch":    {"perc": False, "ref": {"type": "set_min",     "val": 1}},
        "T descarga": {"perc": False, "ref": {"type": "set_max",     "val": 130}},
        "Presion k":  {"perc": False, "ref": {"type": "set_max",     "val": 28}},
        "Presion 0":  {"perc": False, "ref": {"type": "center",      "val": None}},
        "glide k":    {"perc": False, "ref": {"type": "set_max",     "val": 10}},
        "glide 0":    {"perc": False, "ref": {"type": "set_max",     "val": 10}},
    }

    return casos, config_mag


def generar_graficos_ternarios(
    lista_casos: list[dict],
    config: dict[str, dict[str, Any]],
    water_config: str,
) -> None:
    import warnings
    warnings.filterwarnings("ignore", category=UserWarning,
                            message=".*No data for colormapping provided.*")
 
    ASHRAE_NAMES = {
        "DME":       "Dimetil Éter",
        "PROPYLENE": "Propileno",
        "PROPANE":   "Propano",
        "CO2":       "CO2",
        "BUTANE":    "Butano",
        "ISOBUTANE": "Isobutano",
        "METHANE":   "Metano",
        "ETHANE":    "Etano",
        "ETHYLENE":  "Etileno",
    }
 
    def to_ashrae(name: str) -> str:
        return ASHRAE_NAMES.get(name.upper(), name)
 
    def ternary_to_cartesian(a, b, c):
        """Convierte coordenadas ternarias (a, b, c) a cartesianas (x, y).
        a = bottom (eje horizontal), b = right, c = left
        """
        total = a + b + c
        x = 0.5 * (2 * a + b) / total
        y = (np.sqrt(3) / 2) * b / total
        return x, y
 
    def puntos_son_degenerados(x_pts, y_pts, umbral_area=1e-6):
        """Devuelve True si los puntos son colineales o casi colineales
        (el convex hull 2D tiene área menor que umbral_area).
        """
        from scipy.spatial import ConvexHull
        pts_2d = np.stack([x_pts, y_pts], axis=1)
        pts_unicos = np.unique(pts_2d, axis=0)
        if len(pts_unicos) < 3:
            return True
        try:
            hull = ConvexHull(pts_unicos)
            return hull.volume < umbral_area   # en 2D, .volume es el área
        except Exception:
            return True
 
    for caso in tqdm(lista_casos, desc="Generando gráficos ternarios"):
        nombres_lista = caso["nombre"]
        nombre_str    = "_".join(nombres_lista)
        titulo_base   = ", ".join(to_ashrae(n) for n in nombres_lista)
        ejes          = [to_ashrae(n) for n in nombres_lista]
        diccionario_valores = caso["valores"]
 
        carpeta_salida = os.path.join(
            "resultados_ciclo_basico", water_config, "ternarias", "graficos", nombre_str
        )
        os.makedirs(carpeta_salida, exist_ok=True)
 
        for magnitud, conf_data in config.items():
            coords = list(diccionario_valores.keys())
            vals   = [diccionario_valores[c][magnitud] for c in coords]
 
            # ── Guard 1: puntos insuficientes para cualquier interpolación ──
            n_pts = len(coords)
            if n_pts < 3:
                plt.close()
                continue
            # ────────────────────────────────────────────────────────────────
 
            v_min_data, v_max_data = min(vals), max(vals)
 
            ref_type  = conf_data["ref"]["type"]
            ref_val   = conf_data["ref"]["val"]
            perc_mode = conf_data["perc"]
            cmap_name = "coolwarm"
            val_referencia_calculado = None
 
            if ref_type == "set_center":
                center = ref_val
                delta  = max(abs(v_max_data - center), abs(v_min_data - center))
                if delta == 0: delta = 0.001
                vmin, vmax = center - delta, center + delta
                norm = mcolors.TwoSlopeNorm(vmin=vmin, vcenter=center, vmax=vmax)
                val_referencia_calculado = center
 
            elif ref_type == "set_min_max":
                vmin, vmax = ref_val[0], ref_val[1]
                center = (vmin + vmax) / 2
                norm = mcolors.TwoSlopeNorm(vmin=vmin, vcenter=center, vmax=vmax)
                val_referencia_calculado = center
 
            elif ref_type == "set_max":
                vmin, vmax = v_min_data, ref_val
                if vmin >= vmax: vmin = vmax - 0.001
                norm = mcolors.Normalize(vmin=vmin, vmax=vmax)
                cmap_name = "Reds"
                val_referencia_calculado = vmin
 
            elif ref_type == "set_min":
                vmin, vmax = ref_val, v_max_data
                if vmin >= vmax: vmax = vmin + 0.001
                norm = mcolors.Normalize(vmin=vmin, vmax=vmax)
                cmap_name = "Blues_r"
                val_referencia_calculado = vmax
 
            elif ref_type == "center":
                center = (v_min_data + v_max_data) / 2
                delta  = (v_max_data - v_min_data) / 2
                if delta == 0: delta = 0.001
                vmin, vmax = center - delta, center + delta
                norm = mcolors.TwoSlopeNorm(vmin=vmin, vcenter=center, vmax=vmax)
                val_referencia_calculado = center
 
            fig, tax = ternary.figure(scale=1.0)
            fig.set_size_inches(10, 8)
            cmap = plt.get_cmap(cmap_name)
 
            # ── Interpolación y relleno continuo ──────────────────────────────
            coords_arr = np.array(coords)   # shape (N, 3): (a, b, c)
            vals_arr   = np.array(vals)
 
            # Puntos originales en cartesiano
            x_pts, y_pts = ternary_to_cartesian(
                coords_arr[:, 0], coords_arr[:, 1], coords_arr[:, 2]
            )
 
            # Malla densa dentro del triángulo
            grid_res = 300
            a_lin = np.linspace(0, 1, grid_res)
            b_lin = np.linspace(0, 1, grid_res)
            aa, bb = np.meshgrid(a_lin, b_lin)
            cc = 1.0 - aa - bb
            mask = cc >= 0  # solo puntos dentro del triángulo
 
            x_grid, y_grid = ternary_to_cartesian(aa[mask], bb[mask], cc[mask])
 
            pts_origen  = np.stack([x_pts,  y_pts],  axis=1)
            pts_destino = np.stack([x_grid, y_grid], axis=1)
 
            # ── Guard 2: puntos degenerados (colineales o casi colineales) ──
            if puntos_son_degenerados(x_pts, y_pts):
                # nearest no necesita triangulación: nunca lanza QhullError
                try:
                    z_grid = griddata(
                        pts_origen, vals_arr, pts_destino,
                        method="nearest",
                    )
                except Exception:
                    plt.close()
                    continue
            else:
                # ── Flujo normal: cubic → linear → nearest ──────────────────
                try:
                    z_grid = griddata(
                        pts_origen, vals_arr, pts_destino,
                        method="cubic",
                        fill_value=np.nan,
                    )
                    # Si cubic deja demasiados NaN, rellenar con linear
                    nan_ratio = np.isnan(z_grid).sum() / len(z_grid)
                    if nan_ratio > 0.5:
                        z_grid = griddata(
                            pts_origen, vals_arr, pts_destino,
                            method="linear",
                            fill_value=np.nan,
                        )
                except Exception:
                    try:
                        z_grid = griddata(
                            pts_origen, vals_arr, pts_destino,
                            method="linear",
                            fill_value=np.nan,
                        )
                    except Exception:
                        try:
                            z_grid = griddata(
                                pts_origen, vals_arr, pts_destino,
                                method="nearest",
                            )
                        except Exception:
                            plt.close()
                            continue
            # ──────────────────────────────────────────────────────────────────
 
            # Filtrar NaN para tripcolor
            valid = ~np.isnan(z_grid)
            ax = tax.get_axes()
            if valid.sum() >= 3:
                tc = ax.tripcolor(
                    x_grid[valid], y_grid[valid], z_grid[valid],
                    cmap=cmap, norm=norm, shading="gouraud",
                )
                # Contornos suaves encima (opcional, comenta si no los quieres)
                try:
                    ax.tricontour(
                        x_grid[valid], y_grid[valid], z_grid[valid],
                        levels=8, colors="k", linewidths=0.3, alpha=0.3,
                    )
                except Exception:
                    pass
 
            tax.boundary(linewidth=2.0)
            tax.gridlines(color="black", multiple=0.1)
            tax.ticks(axis="lbr", multiple=0.1, linewidth=1, offset=0.02, tick_formats="%.1f")
            tax.get_axes().axis("off")
            tax.clear_matplotlib_ticks()
 
            tax.set_title(f"{titulo_base}: {magnitud}", pad=30, fontsize=15)
            tax.bottom_axis_label(ejes[0], offset=0.06)
            tax.right_axis_label(ejes[1],  offset=0.14)
            tax.left_axis_label(ejes[2],   offset=0.14)
 
            # Colorbar
            if valid.sum() >= 3:
                cb = fig.colorbar(tc, ax=ax, fraction=0.046, pad=0.08)
            else:
                # Fallback: colorbar vacío si no hay datos suficientes
                sm = plt.cm.ScalarMappable(cmap=cmap, norm=norm)
                sm.set_array([])
                cb = fig.colorbar(sm, ax=ax, fraction=0.046, pad=0.08)
 
            ticks_vals = np.linspace(vmin, vmax, 9)
            cb.set_ticks(ticks_vals)
 
            etiquetas = []
            for t in ticks_vals:
                if perc_mode and val_referencia_calculado:
                    variacion = ((t / val_referencia_calculado) - 1) * 100
                    etiquetas.append(f"{'+' if variacion > 0 else ''}{variacion:.1f}%")
                    label_cb = "Variación respecto a Propano"
                else:
                    etiquetas.append(f"{t:.2f}")
                    label_cb = magnitud
 
            cb.set_ticklabels(etiquetas)
            cb.set_label(label_cb)
 
            tax.savefig(os.path.join(carpeta_salida, f"{magnitud}.png"))
            plt.close()

# ---------------------------------------------------------------------------
# Helper de formato Excel
# ---------------------------------------------------------------------------

def _formato_excel(path_excel: str) -> None:
    wb = load_workbook(path_excel)
    ws = wb.active
    align = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = "A2"
    wb.save(path_excel)


# ---------------------------------------------------------------------------
# Gráfico COP/COP_R290 vs VHC/VHC_R290
# ---------------------------------------------------------------------------

# ── Constantes de estilo disponibles ────────────────────────────────────────
# Puedes usar cualquiera de estos en el parámetro `estilo` de cada punto.
#
#  COLORES predefinidos  (también acepta cualquier color válido de matplotlib)
#  ─────────────────────────────────────────────────────────────────────────
#  "azul", "rojo", "verde", "naranja", "morado", "cian", "magenta",
#  "marron", "gris", "negro"
#
#  MARCADORES predefinidos  (también acepta el código directo de matplotlib)
#  ──────────────────────────────────────────────────────────────────────────
#  "circulo"     → "o"     "cuadrado"   → "s"     "triangulo"  → "^"
#  "triangulo_d" → "v"     "rombo"      → "D"     "rombo_p"    → "d"
#  "pentagono"   → "p"     "estrella"   → "*"      "hexagono"   → "h"
#  "cruz"        → "+"     "x"          → "x"     "punto"      → "."
#
#  TAMAÑOS predefinidos  (también acepta un float directamente)
#  ──────────────────────────────────────────────────────────────────────────
#  "pequeño" → 60     "normal" → 120     "grande" → 200     "enorme" → 350
# ────────────────────────────────────────────────────────────────────────────

def grafico_COP_VHC(
    COMPOSICIONES: list[list[float]],
    FLUIDOS: list[list[str]],
    water_config: str = "alta",
    # ── Configuración de cada punto ─────────────────────────────────────────
    estilos: list[dict] | None = None,
    # ── Aspecto general del gráfico ─────────────────────────────────────────
    titulo: str = "Resultados screening, COP vs VHC",
    figsize: tuple[float, float] = (11, 8),
    fuente_anotacion: int = 8,
    # ── Zona de aceptación ──────────────────────────────────────────────────
    margen_vhc: float = 0.30,
    color_zona: str = "steelblue",
    alpha_zona: float = 0.06,
    color_rectas: str = "steelblue",
    lw_rectas: float = 1.6,
    # ── Márgenes del canvas ─────────────────────────────────────────────────
    x_margen: float = 0.15,
    y_margen: float = 0.15,
    # ── Límites manuales de los ejes (None = automático) ────────────────────
    x_min: float | None = None,
    x_max: float | None = None,
    y_min: float | None = None,
    y_max: float | None = None,
    # ── Salida ──────────────────────────────────────────────────────────────
    guardar_ruta: str | None = None,
    dpi: int = 150,
    mostrar: bool = False,
) -> plt.Figure:
    """
    Genera un gráfico de dispersión con:
      • Eje X : VHC / VHC_{R290}
      • Eje Y : COP / COP_{R290}
 
    Los nombres de fluidos se muestran en notación ASHRAE cuando están
    disponibles en el diccionario interno ASHRAE_NAMES. Los cálculos siguen
    usando los nombres originales de REFPROP.
 
    Las tres rectas delimitan la zona de aceptación (relleno azul tenue):
      • x = 1 - margen_vhc  (VHC mínimo aceptable)
      • x = 1 + margen_vhc  (VHC máximo aceptable)
      • y = 1               (COP al menos igual al propano)
    """
 
    # ── Traducción de nombres REFPROP → ASHRAE ────────────────────────────────
    ASHRAE_NAMES = {
        "DME":       "RE170",
        "PROPYLENE": "R1270",
        "PROPANE":   "R290",
        "CO2":       "CO2",
        "BUTANE":    "R600",
        "ISOBUTANE": "R600a",
        "METHANE":   "R50",
        "ETHANE":    "R170",
        "ETHYLENE":  "R1150",
    }
 
    def _to_ashrae(nombre: str) -> str:
        """Devuelve el nombre ASHRAE si existe, o el original en su defecto."""
        return ASHRAE_NAMES.get(nombre.upper(), nombre)
 
    # ── Tablas de traducción de alias ─────────────────────────────────────────
    _COLORES = {
        "azul": "steelblue", "rojo": "firebrick", "verde": "seagreen",
        "naranja": "darkorange", "morado": "mediumpurple", "cian": "darkcyan",
        "magenta": "deeppink", "marron": "saddlebrown", "gris": "slategray",
        "negro": "black",
    }
    _MARCADORES = {
        "circulo": "o", "cuadrado": "s", "triangulo": "^", "triangulo_d": "v",
        "rombo": "D", "rombo_p": "d", "pentagono": "p", "estrella": "*",
        "hexagono": "h", "cruz": "+", "x": "x", "punto": ".",
    }
    _TAMANIOS = {
        "pequeño": 60, "normal": 120, "grande": 200, "enorme": 350,
    }
 
    def _resolver_estilo(d: dict) -> dict:
        """Resuelve alias y aplica defaults para un dict de estilo."""
        color    = d.get("color",    "steelblue")
        marcador = d.get("marcador", "o")
        tamanio  = d.get("tamanio",  "normal")
        alpha    = d.get("alpha",    0.85)
        offset   = d.get("etiqueta_offset", (0, 0.02))
        anotacion = d.get("anotacion", None)
        return {
            "color":    _COLORES.get(color, color),
            "marcador": _MARCADORES.get(marcador, marcador),
            "tamanio":  _TAMANIOS.get(tamanio, tamanio) if isinstance(tamanio, str) else tamanio,
            "alpha":    alpha,
            "offset":   offset,
            "anotacion": anotacion,
        }
 
    # ── Valores de referencia R290 ────────────────────────────────────────────
    vhc_min_ref, vhc_max_ref, cop_ref = calcular_valores_referencia(water_config)
    vhc_r290 = vhc_min_ref / (1 - margen_vhc)
    cop_r290 = cop_ref
 
    # ── Calcular COP y VHC de cada punto de entrada ───────────────────────────
    xs: list[float] = []
    ys: list[float] = []
    etiquetas_fluidos: list[str] = []   # nombres ASHRAE para mostrar
    etiquetas_comps:   list[str] = []
    errores: list[str | None] = []
 
    for fluidos, comps in zip(FLUIDOS, COMPOSICIONES):
        n = len(fluidos)
 
        # Normalizar a ternaria para REFPROP
        if n == 1:
            fA, fB, fC = fluidos[0], fluidos[0], fluidos[0]
            xA, xB = 1.0, 0.0
        elif n == 2:
            fA, fB, fC = fluidos[0], fluidos[1], fluidos[1]
            xA, xB = comps[0], comps[1]
        else:
            fA, fB, fC = fluidos[0], fluidos[1], fluidos[2]
            xA, xB = comps[0], comps[1]
 
        # Cálculo con nombres originales REFPROP
        row = calcular_ciclo_basico(fA, fB, fC, xA, xB, water_config)
 
        err = row.get("error")
        errores.append(err)
 
        if err:
            xs.append(float("nan"))
            ys.append(float("nan"))
        else:
            xs.append(row["VHC"] / vhc_r290)
            ys.append(row["COP"] / cop_r290)
 
        # Etiquetas con nombres ASHRAE
        nombres_ashrae = [_to_ashrae(f) for f in fluidos]
        etiquetas_fluidos.append(", ".join(nombres_ashrae))
        etiquetas_comps.append(
            None if n == 1
            else ", ".join(f"{c*100:.0f}%" for c in comps)
        )
 
    # ── Construir figura ──────────────────────────────────────────────────────
    fig, ax = plt.subplots(figsize=figsize)
 
    # ── Determinar rango del canvas ───────────────────────────────────────────
    xs_validos = [v for v in xs if not np.isnan(v)]
    ys_validos = [v for v in ys if not np.isnan(v)]
 
    x_lo = x_min if x_min is not None else (min(xs_validos + [1 - margen_vhc], default=0.5) - x_margen)
    x_hi = x_max if x_max is not None else (max(xs_validos + [1 + margen_vhc], default=1.5) + x_margen)
    y_lo = y_min if y_min is not None else (min(ys_validos + [1.0], default=0.5) - y_margen)
    y_hi = y_max if y_max is not None else (max(ys_validos + [1.0], default=1.5) + y_margen)
 
    # ── Zona de aceptación: relleno azul tenue ────────────────────────────────
    x_izq = 1 - margen_vhc
    x_der = 1 + margen_vhc
    y_base = 1.0
    y_hi_seg = y_hi + 0.5
 
    # Relleno interior de la zona de aceptación
    ax.fill_betweenx(
        [y_base, y_hi_seg],
        x_izq, x_der,
        color=color_zona,
        alpha=alpha_zona,
        zorder=1,
    )
 
    # Tres segmentos delimitadores (U abierta hacia arriba)
    lineas_kwargs = dict(color=color_rectas, lw=lw_rectas, ls="-", zorder=2)
    ax.plot([x_izq, x_izq], [y_base, y_hi_seg], **lineas_kwargs, label="_nolegend_")
    ax.plot([x_izq, x_der], [y_base, y_base],   **lineas_kwargs, label="_nolegend_")
    ax.plot([x_der, x_der], [y_base, y_hi_seg], **lineas_kwargs, label="_nolegend_")
 
    # ── Punto de referencia R290 ──────────────────────────────────────────────
    ax.scatter(1.0, 1.0, color="black", marker="*", s=220, zorder=5)
    ax.annotate(
        "R290",
        xy=(1.0, 1.0),
        xytext=(1.0, 1.0 - 0.018),
        fontsize=fuente_anotacion,
        color="black",
        fontweight="bold",
        ha="center",
    )
 
    # ── Plotear cada punto ────────────────────────────────────────────────────
    estilos = estilos or []
 
    for i, (x, y) in enumerate(zip(xs, ys)):
        est_raw = estilos[i] if i < len(estilos) else {}
        est = _resolver_estilo(est_raw)
 
        if np.isnan(x) or np.isnan(y):
            print(
                f"[grafico_COP_VHC] Punto {i} ({etiquetas_fluidos[i]}) "
                f"no calculado: {errores[i]}"
            )
            continue
 
        ax.scatter(
            x, y,
            c=est["color"],
            marker=est["marcador"],
            s=est["tamanio"],
            alpha=est["alpha"],
            edgecolors="white",
            linewidths=0.6,
            zorder=6,
        )
 
        # Anotación con nombres ASHRAE
        dx, dy = est["offset"]
        if etiquetas_comps[i] is None:
            texto = etiquetas_fluidos[i]
        else:
            texto = f"{etiquetas_fluidos[i]}\n{etiquetas_comps[i]}"
        if est["anotacion"]:
            texto += f"\n{est['anotacion']}"
        ax.annotate(
            texto,
            xy=(x, y),
            xytext=(x + dx, y + dy),
            fontsize=fuente_anotacion,
            color=est["color"],
            ha="center",
            va="bottom",
            bbox=dict(
                boxstyle="round,pad=0.25",
                facecolor="white",
                edgecolor=est["color"],
                alpha=0.75,
                linewidth=0.8,
            ),
            arrowprops=dict(
                arrowstyle="-",
                color=est["color"],
                lw=0.7,
                alpha=0.6,
            ),
        )
 
    # ── Ejes y estética ───────────────────────────────────────────────────────
    ax.set_xlim(x_lo, x_hi)
    ax.set_ylim(y_lo, y_hi)
 
    ax.set_xlabel(r"$\mathrm{VHC}\,/\,\mathrm{VHC}_{R290}$", fontsize=12)
    ax.set_ylabel(r"$\mathrm{COP}\,/\,\mathrm{COP}_{R290}$",  fontsize=12)
    ax.set_title(titulo, fontsize=13, pad=12)
 
    ax.grid(True, linestyle=":", linewidth=0.6, alpha=0.5)
    ax.set_axisbelow(True)
 
    fig.tight_layout()
 
    if guardar_ruta:
        os.makedirs(os.path.dirname(guardar_ruta) or ".", exist_ok=True)
        fig.savefig(guardar_ruta, dpi=dpi, bbox_inches="tight")
        print(f"[grafico_COP_VHC] Figura guardada en: {guardar_ruta}")
 
    if mostrar:
        plt.show()
 
    return fig


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    init_refprop()

    water_config = "alta"  # media, alta, media_7, alta_7
    posibles_refrigerantes = ["PROPANE", "PROPYLENE", "DME", "CO2", "ETHANE", "ETHYLENE", "BUTANE", "ISOBUTANE", "METHANE"]
    n_prop = 21  # 5% de salto entre proporción y proporción de refrigerante

    # # 1. Cálculo bruto
    # calcular_resultados(posibles_refrigerantes, water_config, n_prop)
    # df_a_excel(water_config)

    # # 2. Filtrado bruto
    # df_a_excel_filtrado(water_config)

    # # 3. Refinado fino
    # refinar_mezclas(water_config)
    # df_a_excel_fino(water_config)

    # # 4. Resumen y TXT
    # crear_excel_resumen(water_config)
    # guardar_txt(water_config)

    # # 5. Gráficos ternarios
    # casos, config_mag = obtener_casos(water_config)
    # generar_graficos_ternarios(casos, config_mag, water_config)





# ======================
# TEMPERATURA MEDIA
# ======================
#     grafico_COP_VHC(
#     water_config=water_config,
#     COMPOSICIONES = [
#         [0.03, 0.66, 0.31],
#         [0.15, 0.51, 0.34],
#         [0.19, 0.65, 0.16],
#         [1],
#         [0.7, 0.3],
#         [0.6, 0.2, 0.2],
#         [0.5, 0.5],
#         [0.9, 0.1],
#         [0.55, 0.45],
#         [0.6, 0.4],
#         [0.85, 0.15]
#     ],
#     FLUIDOS = [
#         ["PROPYLENE", "DME", "BUTANE"],
#         ["PROPANE", "DME", "BUTANE"],
#         ["PROPYLENE", "DME", "ISOBUTANE"],
#         ["DME"],
#         ["PROPANE", "BUTANE"],
#         ["PROPYLENE", "BUTANE", "ISOBUTANE"],
#         ["PROPYLENE", "ISOBUTANE"],
#         ["PROPYLENE", "BUTANE"],
#         ["PROPYLENE", "ISOBUTANE"],
#         ["PROPANE", "ISOBUTANE"],
#         ["PROPANE", "DME"]

#     ],
#     estilos = [
#         {"color": "azul",    "marcador": "circulo",         "etiqueta_offset": (0, 0.015)},
#         {"color": "rojo", "marcador": "cuadrado",           "etiqueta_offset": (0, 0.01)},
#         {"color": "verde",     "marcador": "rombo"    ,     "etiqueta_offset": (0, 0.015)},
#         {"color": "naranja",     "marcador": "triangulo"  , "etiqueta_offset": (0, 0.015)},
#         {"color": "morado",     "marcador": "triangulo_d" , "etiqueta_offset": (0.02, 0.015)},
#         {"color": "cian",     "marcador": "rombo_p"    ,    "etiqueta_offset": (0.06, 0)},
#         {"color": "magenta",     "marcador": "pentagono" ,  "etiqueta_offset": (0, 0.02)},
#         {"color": "marron",     "marcador": "estrella" ,    "etiqueta_offset": (0, 0.015)},
#         {"color": "gris",     "marcador": "hexagono" ,      "etiqueta_offset": (0.045, -0.015)},
#         {"color": "negro",     "marcador": "x" ,            "etiqueta_offset": (0.02, -0.025)},
#         {"color": "rojo",     "marcador": "rombo" ,   "etiqueta_offset": (0, 0.015), "anotacion": "Erik Mickoleit et al."}
#     ],
#     margen_vhc   = 0.30,
#     guardar_ruta = "resultados/mapa_COP_VHC.png",

#     y_min = 0.95,
#     y_max = 1.20,
#     x_min = 0.65,
#     x_max = 1.35
# )

# ==========================
# TEMPERATURA ALTA
# ==========================

#     grafico_COP_VHC(
#     water_config=water_config,
#     COMPOSICIONES = [
#         [0.04, 0.72, 0.24],
#         [0.15, 0.52, 0.33],
#         [0.19, 0.65, 0.16],
#         [0.99, 0.01],
#         [0.93, 0.07],
#         [0.95, 0.05],
#         [0.98, 0.02],
#         [0.99, 0.01],
#         [0.99, 0.01],
#         [1],
#         [0.85, 0.15]
#     ],
#     FLUIDOS = [
#         ["PROPYLENE", "DME", "BUTANE"],
#         ["PROPANE", "DME", "BUTANE"],
#         ["PROPYLENE", "DME", "ISOBUTANE"],
#         ["DME", "CO2"],
#         ["DME", "PROPYLENE"],
#         ["DME", "PROPANE"],
#         ["DME", "ETHANE"],
#         ["DME", "ETHYLENE"],
#         ["DME", "BUTANE"],
#         ["DME"],
#         ["PROPANE", "DME"]

#     ],
#     estilos = [
#         {"color": "azul",    "marcador": "circulo",         "etiqueta_offset": (0, 0.02)},
#         {"color": "rojo", "marcador": "cuadrado",           "etiqueta_offset": (0, 0.01)},
#         {"color": "verde",     "marcador": "rombo"    ,     "etiqueta_offset": (0, 0.02)},
#         {"color": "naranja",     "marcador": "triangulo"  , "etiqueta_offset": (-0.01, 0.01)},
#         {"color": "morado",     "marcador": "triangulo_d" , "etiqueta_offset": (0.04, 0)},
#         {"color": "cian",     "marcador": "rombo_p"    ,    "etiqueta_offset": (0.03, -0.015)},
#         {"color": "magenta",     "marcador": "pentagono" ,  "etiqueta_offset": (0.04, 0.02)},
#         {"color": "marron",     "marcador": "estrella" ,    "etiqueta_offset": (0.015, -0.03)},
#         {"color": "gris",     "marcador": "hexagono" ,      "etiqueta_offset": (-0.045, 0.01)},
#         {"color": "negro",     "marcador": "x" ,            "etiqueta_offset": (-0.02, -0.02)},
#         {"color": "rojo",     "marcador": "rombo" ,   "etiqueta_offset": (0, 0.015), "anotacion": "Erik Mickoleit et al."}
#     ],
#     margen_vhc   = 0.30,
#     guardar_ruta = "resultados/mapa_COP_VHC.png",

#     y_min = 0.95,
#     y_max = 1.20,
#     x_min = 0.65,
#     x_max = 1.35
# )


if __name__ == "__main__":
    main()