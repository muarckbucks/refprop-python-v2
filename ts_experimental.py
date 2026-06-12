"""
plot_ts_diagrams.py
-------------------
Genera un diagrama T-S comparativo para múltiples fluidos/mezclas en una
sola figura de matplotlib. No abre la ventana; devuelve el objeto Figure.

rprop trabaja en °C y bar.

Dependencias: matplotlib, numpy, openpyxl, ciclo_ternario (rprop,
              calcular_ciclo_basico, init_refprop).

Arquitectura
------------
Cada entrada al gráfico es un objeto `FluidSeries`, que agrupa:

  * el fluido (lista de componentes) y su composición molar
  * opciones visuales (color, grosor de línea, etc.)
  * opcionalmente, un ciclo termodinámico cuya fuente puede ser:
      - CycleSource.THEORETICAL  → calculado con REFPROP / calcular_ciclo_basico
      - CycleSource.EXPERIMENTAL → puntos leídos desde un archivo .xlsx
      - CycleSource.BOTH         → ambos superpuestos

Los datos experimentales (CycleDef con source != THEORETICAL) se configuran
directamente dentro de FluidSeries, en lugar de ir en una lista separada.
Esto permite que la leyenda, los colores y la lógica de dibujado sean
coherentes para cada fluido.
"""

from __future__ import annotations

import numpy as np
import matplotlib
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.colors import to_rgba
from dataclasses import dataclass, field
from enum import Enum
from typing import Optional

import openpyxl

from ciclo_ternario import rprop, calcular_ciclo_basico, init_refprop

matplotlib.use("Agg")  # backend sin pantalla


# ---------------------------------------------------------------------------
# Nombres ASHRAE
# ---------------------------------------------------------------------------

ASHRAE_NAMES = {
    "DME":       "RE170",
    "PROPYLENE": "R1270",
    "PROPANE":   "R290",
    "CO2":       "CO2",
    "BUTANE":    "R600",
    "ISOBUTANE": "R600a",
    "METHANE":   "R50",
    "ETHANE":    "R170",
    "ETHYLENE":  "R1150",
}

# Presión mínima de saturación [bar] por fluido (usado para acotar T_min).
_P_MIN_SAT: dict[str, float] = {
    "CO2": 5.5,
}


# ---------------------------------------------------------------------------
# Fuente del ciclo
# ---------------------------------------------------------------------------

class CycleSource(Enum):
    """
    Indica de dónde provienen los puntos del ciclo termodinámico.

    THEORETICAL  – calculado con REFPROP (calcular_ciclo_basico).
    EXPERIMENTAL – leído desde un archivo .xlsx.
    BOTH         – ambos superpuestos en el mismo color.
    """
    THEORETICAL  = "theoretical"
    EXPERIMENTAL = "experimental"
    BOTH         = "both"


# ---------------------------------------------------------------------------
# Definición del ciclo (teórico + experimental en un solo objeto)
# ---------------------------------------------------------------------------

@dataclass
class CycleDef:
    """
    Define el ciclo a superponer sobre la curva de saturación de un FluidSeries.

    Parámetros comunes
    ------------------
    source : CycleSource
        Origen de los datos: THEORETICAL, EXPERIMENTAL o BOTH.
    linewidth : float
        Grosor de la línea del ciclo (default 1.5).
    linestyle : str
        Estilo de la línea del ciclo (default "-").
    alpha : float
        Opacidad de la línea del ciclo (default 0.55).
    marker : str
        Marcador para los puntos del ciclo (default "o").
    markersize : float
        Tamaño del marcador (default 7).

    Parámetros para source = THEORETICAL o BOTH
    --------------------------------------------
    water_config : str
        Configuración del agua, pasada a calcular_ciclo_basico (default "default").

    Parámetros para source = EXPERIMENTAL o BOTH
    ---------------------------------------------
    xlsx_path : str
        Ruta al archivo Excel (.xlsx).
    sheet_name : str
        Nombre de la hoja dentro del libro.
    row : int
        Número de fila (base 1) donde se encuentran los datos del ensayo.
    temperature_cols : dict[str, str]
        Columnas de temperatura [°C]. Ej: {"T1": "B", "T2": "C", "T3": "D", "T4": "E"}
    entropy_cols : dict[str, str]
        Columnas de entropía [kJ/(kg·K)]. Ej: {"s1": "F", "s2": "G", "s3": "H", "s4": "I"}
    exp_linestyle : str
        Estilo de línea específico para el ciclo experimental cuando source=BOTH
        (default "--"). Si source=EXPERIMENTAL, se usa `linestyle`.
    exp_alpha : float
        Opacidad específica para el ciclo experimental cuando source=BOTH
        (default 0.7). Si source=EXPERIMENTAL, se usa `alpha`.

    Notas
    -----
    * Si source es THEORETICAL, los parámetros experimentales se ignoran.
    * Si source es EXPERIMENTAL, los parámetros teóricos se ignoran.
    * Si source es BOTH, el ciclo teórico usa (linestyle, alpha) y el
      experimental usa (exp_linestyle, exp_alpha).

    Ejemplo — solo teórico
    ----------------------
    >>> CycleDef(source=CycleSource.THEORETICAL, water_config="media")

    Ejemplo — solo experimental
    ---------------------------
    >>> CycleDef(
    ...     source           = CycleSource.EXPERIMENTAL,
    ...     xlsx_path        = "datos.xlsx",
    ...     sheet_name       = "Ensayos",
    ...     row              = 5,
    ...     temperature_cols = {"T1": "B", "T2": "C", "T3": "D", "T4": "E"},
    ...     entropy_cols     = {"s1": "F", "s2": "G", "s3": "H", "s4": "I"},
    ... )

    Ejemplo — ambos
    ---------------
    >>> CycleDef(
    ...     source           = CycleSource.BOTH,
    ...     water_config     = "media",
    ...     xlsx_path        = "datos.xlsx",
    ...     sheet_name       = "Ensayos",
    ...     row              = 5,
    ...     temperature_cols = {"T1": "B", "T2": "C", "T3": "D", "T4": "E"},
    ...     entropy_cols     = {"s1": "F", "s2": "G", "s3": "H", "s4": "I"},
    ...     exp_linestyle    = "--",
    ...     exp_alpha        = 0.7,
    ... )
    """
    source: CycleSource = CycleSource.THEORETICAL

    # ── apariencia común ────────────────────────────────────────────────────
    linewidth:   float = 1.5
    linestyle:   str   = "-"
    alpha:       float = 0.55
    marker:      str   = "o"
    markersize:  float = 7

    # ── parámetros teóricos ─────────────────────────────────────────────────
    water_config: str = "default"

    # ── parámetros experimentales ───────────────────────────────────────────
    xlsx_path:        Optional[str]             = None
    sheet_name:       Optional[str]             = None
    row:              Optional[int]             = None
    temperature_cols: Optional[dict[str, str]]  = None
    entropy_cols:     Optional[dict[str, str]]  = None

    # ── apariencia diferenciada para BOTH ───────────────────────────────────
    exp_linestyle: str   = "--"
    exp_alpha:     float = 0.7

    def _requires_experimental(self) -> bool:
        return self.source in (CycleSource.EXPERIMENTAL, CycleSource.BOTH)

    def _requires_theoretical(self) -> bool:
        return self.source in (CycleSource.THEORETICAL, CycleSource.BOTH)

    def validate(self) -> None:
        """Lanza ValueError si faltan parámetros obligatorios."""
        if self._requires_experimental():
            missing = [
                name for name, val in [
                    ("xlsx_path",        self.xlsx_path),
                    ("sheet_name",       self.sheet_name),
                    ("row",              self.row),
                    ("temperature_cols", self.temperature_cols),
                    ("entropy_cols",     self.entropy_cols),
                ]
                if val is None
            ]
            if missing:
                raise ValueError(
                    f"CycleDef con source={self.source.value} requiere: "
                    + ", ".join(missing)
                )


# ---------------------------------------------------------------------------
# Serie de fluido (unidad principal de configuración)
# ---------------------------------------------------------------------------

@dataclass
class FluidSeries:
    """
    Agrupa toda la información de un fluido/mezcla para el diagrama T-S.

    Parámetros obligatorios
    -----------------------
    fluid : list[str]
        Nombres de los componentes según la DLL de REFPROP.
        Ej: ["PROPANE"] o ["PROPYLENE", "DME"].
    composition : list[float]
        Fracciones molares. Deben sumar 1.
        Ej: [1.0] o [0.20, 0.80].

    Parámetros opcionales
    ---------------------
    label : str | None
        Etiqueta en la leyenda. None → generada automáticamente con nombres ASHRAE.
    color : str | None
        Color matplotlib. None → asignado automáticamente desde la paleta tab10.
    linewidth : float
        Grosor de la curva de saturación (default 2.0).
    linestyle : str
        Estilo de la curva de saturación (default "-").
    n_points : int
        Resolución de la curva de saturación (default 200).
    cycle : CycleDef | None
        Define el ciclo a superponer. None → no se dibuja ningún ciclo.

    Ejemplo — curva de saturación sin ciclo
    ----------------------------------------
    >>> FluidSeries(fluid=["PROPANE"], composition=[1.0])

    Ejemplo — con ciclo teórico
    ---------------------------
    >>> FluidSeries(
    ...     fluid       = ["PROPANE"],
    ...     composition = [1.0],
    ...     color       = "steelblue",
    ...     cycle       = CycleDef(source=CycleSource.THEORETICAL,
    ...                            water_config="media"),
    ... )

    Ejemplo — con ciclo experimental
    ---------------------------------
    >>> FluidSeries(
    ...     fluid       = ["PROPYLENE", "DME"],
    ...     composition = [0.20, 0.80],
    ...     color       = "tomato",
    ...     cycle       = CycleDef(
    ...         source           = CycleSource.EXPERIMENTAL,
    ...         xlsx_path        = "datos.xlsx",
    ...         sheet_name       = "Ensayos",
    ...         row              = 4,
    ...         temperature_cols = {"T1": "B", "T2": "C", "T3": "D", "T4": "E"},
    ...         entropy_cols     = {"s1": "F", "s2": "G", "s3": "H", "s4": "I"},
    ...         linestyle        = "--",
    ...         alpha            = 0.7,
    ...     ),
    ... )

    Ejemplo — con ambos ciclos
    --------------------------
    >>> FluidSeries(
    ...     fluid       = ["PROPANE"],
    ...     composition = [1.0],
    ...     color       = "steelblue",
    ...     cycle       = CycleDef(
    ...         source           = CycleSource.BOTH,
    ...         water_config     = "media",
    ...         xlsx_path        = "datos.xlsx",
    ...         sheet_name       = "Ensayos",
    ...         row              = 4,
    ...         temperature_cols = {"T1": "B", "T2": "C", "T3": "D", "T4": "E"},
    ...         entropy_cols     = {"s1": "F", "s2": "G", "s3": "H", "s4": "I"},
    ...     ),
    ... )
    """
    fluid:       list[str]
    composition: list[float]

    label:     Optional[str]   = None
    color:     Optional[str]   = None
    linewidth: float           = 2.0
    linestyle: str             = "-"
    n_points:  int             = 200

    cycle:     Optional[CycleDef] = None


# ---------------------------------------------------------------------------
# Helpers internos
# ---------------------------------------------------------------------------

def _ashrae_label(fluid: list[str], comp: list[float]) -> str:
    """
    Genera la etiqueta de leyenda con nombres ASHRAE.
    'RE170' para fluido puro; 'R290/R1270 (20%/80%)' para mezclas.
    """
    names = [ASHRAE_NAMES.get(f, f) for f in fluid]
    if len(fluid) == 1:
        return names[0]
    name_str = "/".join(names)
    frac_str  = "/".join(f"{round(c*100):g}%" for c in comp)
    return f"{name_str} ({frac_str})"


def _sat_curve(
    fluid: list[str],
    composition: list[float],
    n_points: int = 200,
    p_min_sat: Optional[float] = None,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """
    Calcula la curva de saturación (líquido + vapor).
    Devuelve s_liq, T_liq, s_vap, T_vap (arrays numpy).

    Eje X → entropía específica s [kJ/(kg·K)]
    Eje Y → temperatura T [°C]
    """
    T_crit_C = rprop(fluid, "Tcrit", composition, T=0, H=0)

    T_min_C: Optional[float] = None
    for T_probe in np.arange(-180, T_crit_C - 1, 1.0):
        try:
            rprop(fluid, "S", composition, T=float(T_probe), Q=0)
            T_min_C = float(T_probe)
            break
        except Exception:
            continue

    if T_min_C is None:
        raise RuntimeError(
            f"No se encontró T_min para {fluid} entre -180 °C y {T_crit_C:.1f} °C"
        )

    if p_min_sat is not None:
        for T_probe in np.arange(T_min_C, T_crit_C - 1, 0.5):
            try:
                p_test = rprop(fluid, "P", composition, T=float(T_probe), Q=0)
                if p_test >= p_min_sat:
                    T_min_C = float(T_probe)
                    break
            except Exception:
                continue

    temps = np.linspace(T_min_C, T_crit_C - 0.05, n_points)
    s_liq, T_liq, s_vap, T_vap = [], [], [], []

    for T in temps:
        try:
            s_l = rprop(fluid, "S", composition, T=float(T), Q=0)
            s_v = rprop(fluid, "S", composition, T=float(T), Q=1)
            s_liq.append(s_l); T_liq.append(T)
            s_vap.append(s_v); T_vap.append(T)
        except Exception:
            continue

    return (np.array(s_liq), np.array(T_liq),
            np.array(s_vap), np.array(T_vap))


# ── Ciclo teórico ────────────────────────────────────────────────────────────

def _compute_theoretical_cycle(
    fluid: list[str],
    composition: list[float],
    water_config: str,
) -> Optional[dict]:
    """
    Llama a calcular_ciclo_basico y devuelve el dict de resultados,
    o None si falla.
    """
    f_pad = (fluid + ["", ""])[:3]
    c_pad = (composition + [0.0, 0.0])[:3]
    try:
        row = calcular_ciclo_basico(
            fluid_A=f_pad[0], fluid_B=f_pad[1], fluid_C=f_pad[2],
            x_A=c_pad[0], x_B=c_pad[1],
            water_config=water_config,
        )
        return row
    except Exception as exc:
        print(f"[AVISO] No se pudo calcular el ciclo teórico para {fluid}: {exc}")
        return None


def _draw_theoretical_cycle(
    ax,
    cycle_row: dict,
    color,
    linewidth: float,
    linestyle: str,
    alpha: float,
) -> None:
    """
    Dibuja las líneas 1→2→3→4→1 del ciclo teórico en el plano T-S.
    Requiere las claves s_1…s_4 y T_1…T_4 en cycle_row.
    """
    if cycle_row.get("error") is not None:
        return
    try:
        s1, T1 = cycle_row["s_1"], cycle_row["T_1"]
        s2, T2 = cycle_row["s_2"], cycle_row["T_2"]
        s3, T3 = cycle_row["s_3"], cycle_row["T_3"]
        s4, T4 = cycle_row["s_4"], cycle_row["T_4"]
    except KeyError:
        return

    c = to_rgba(color, alpha=alpha)
    kw = dict(color=c, linewidth=linewidth, linestyle=linestyle,
              solid_capstyle="round", solid_joinstyle="round")
    ax.plot([s1, s2], [T1, T2], **kw)
    ax.plot([s2, s3], [T2, T3], **kw)
    ax.plot([s3, s4], [T3, T4], **kw)
    ax.plot([s4, s1], [T4, T1], **kw)


# ── Ciclo experimental ───────────────────────────────────────────────────────

def _col_to_index(col: str | int) -> int:
    if isinstance(col, int):
        return col
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(str(col).strip().upper())


def _load_experimental_points(
    cycle_def: CycleDef,
    series_label: str,
) -> tuple[dict[str, float], dict[str, float]]:
    """
    Lee temperaturas [°C] y entropías [kJ/(kg·K)] desde el .xlsx definido
    en cycle_def. Devuelve (temperatures, entropies).
    Celdas vacías o no numéricas se omiten con aviso.
    """
    wb = openpyxl.load_workbook(cycle_def.xlsx_path, data_only=True, read_only=True)

    if cycle_def.sheet_name not in wb.sheetnames:
        raise ValueError(
            f"La hoja '{cycle_def.sheet_name}' no existe en '{cycle_def.xlsx_path}'. "
            f"Hojas disponibles: {wb.sheetnames}"
        )
    ws = wb[cycle_def.sheet_name]

    temperatures: dict[str, float] = {}
    entropies:    dict[str, float] = {}

    def _read_cell(key: str, col: str | int, dest: dict) -> None:
        col_idx = _col_to_index(col)
        cell    = ws.cell(row=cycle_def.row, column=col_idx)
        val     = cell.value
        if val is None:
            print(
                f"[AVISO] '{series_label}' (experimental): "
                f"celda ({cycle_def.row}, {col}) → '{key}' está vacía. Se omite."
            )
            return
        try:
            dest[key] = float(val)
        except (TypeError, ValueError):
            print(
                f"[AVISO] '{series_label}' (experimental): "
                f"celda ({cycle_def.row}, {col}) → '{key}' = {val!r} no es numérico. "
                f"Se omite."
            )

    for key, col in cycle_def.temperature_cols.items():
        _read_cell(key, col, temperatures)
    for key, col in cycle_def.entropy_cols.items():
        _read_cell(key, col, entropies)

    wb.close()
    return temperatures, entropies


def _draw_experimental_cycle(
    ax,
    color,
    temperatures: dict[str, float],
    entropies: dict[str, float],
    marker: str,
    markersize: float,
    linewidth: float,
    linestyle: str,
    alpha: float,
    series_label: str,
    zorder: int = 5,
) -> None:
    """
    Dibuja los puntos experimentales y, si hay más de uno, el ciclo cerrado
    en el plano T-S. Empareja "sN" con "TN" por sufijo.
    """
    paired_s: list[float] = []
    paired_T: list[float] = []
    paired_idx: list[int] = []

    for s_key, s_val in entropies.items():
        T_key = "T" + s_key[1:]
        if T_key not in temperatures:
            print(
                f"[AVISO] '{series_label}' (experimental): "
                f"no se encontró temperatura para '{s_key}' (esperada '{T_key}'). "
                f"Punto omitido."
            )
            continue
        paired_s.append(s_val)
        paired_T.append(temperatures[T_key])
        suffix = "".join(c for c in s_key if c.isdigit())
        paired_idx.append(int(suffix) if suffix else 0)

    if not paired_s:
        print(f"[AVISO] '{series_label}' (experimental): ningún punto válido. Se omite.")
        return

    ax.scatter(
        paired_s, paired_T,
        color=color, marker=marker, s=markersize ** 2, zorder=zorder,
    )

    if len(paired_s) > 1:
        order   = sorted(range(len(paired_idx)), key=lambda i: paired_idx[i])
        s_ord   = [paired_s[i] for i in order]
        T_ord   = [paired_T[i] for i in order]
        s_closed = s_ord + [s_ord[0]]
        T_closed = T_ord + [T_ord[0]]

        c = to_rgba(color, alpha=alpha)
        ax.plot(
            s_closed, T_closed,
            color=c, linewidth=linewidth, linestyle=linestyle,
            solid_capstyle="round", solid_joinstyle="round",
            zorder=zorder - 1,
        )


# ---------------------------------------------------------------------------
# Función principal
# ---------------------------------------------------------------------------

def plot_ts_diagrams(
    series: list[FluidSeries],
    *,
    # ── límites manuales del gráfico ────────────────────────────────────────
    s_min: Optional[float] = None,
    s_max: Optional[float] = None,
    T_min: Optional[float] = None,
    T_max: Optional[float] = None,
    # ── opciones de figura ──────────────────────────────────────────────────
    figsize: tuple[float, float] = (12, 7),
    dpi:     int   = 150,
    title:   str   = "Diagramas T-S comparativos",
    xlabel:  str   = "Entropía específica  s  [kJ/(kg·K)]",
    ylabel:  str   = "Temperatura  T  [°C]",
    grid:    bool  = True,
) -> plt.Figure:
    """
    Dibuja las curvas de saturación de varias series en un solo diagrama T-S.

    Parámetros
    ----------
    series : list[FluidSeries]
        Lista de fluidos/mezclas a representar. Cada elemento es un
        `FluidSeries` que incluye, opcionalmente, un `CycleDef`.

    s_min, s_max : float | None
        Límites eje X [kJ/(kg·K)].
    T_min, T_max : float | None
        Límites eje Y [°C].
    figsize, dpi, title, xlabel, ylabel, grid
        Opciones estándar de figura matplotlib.

    Devuelve
    --------
    fig : matplotlib.figure.Figure
        Figura sin mostrar. Úsala con fig.savefig(...) o fig.show().

    Ejemplos
    --------
    # Solo curvas de saturación
    >>> fig = plot_ts_diagrams([
    ...     FluidSeries(fluid=["PROPANE"],            composition=[1.0]),
    ...     FluidSeries(fluid=["PROPYLENE", "DME"],   composition=[0.20, 0.80]),
    ... ])

    # Con ciclos: uno teórico y uno experimental
    >>> fig = plot_ts_diagrams([
    ...     FluidSeries(
    ...         fluid=["PROPANE"], composition=[1.0], color="steelblue",
    ...         cycle=CycleDef(source=CycleSource.THEORETICAL, water_config="media"),
    ...     ),
    ...     FluidSeries(
    ...         fluid=["PROPYLENE", "DME"], composition=[0.20, 0.80],
    ...         color="tomato",
    ...         cycle=CycleDef(
    ...             source           = CycleSource.EXPERIMENTAL,
    ...             xlsx_path        = "datos.xlsx",
    ...             sheet_name       = "Ensayos",
    ...             row              = 4,
    ...             temperature_cols = {"T1":"B","T2":"C","T3":"D","T4":"E"},
    ...             entropy_cols     = {"s1":"F","s2":"G","s3":"H","s4":"I"},
    ...         ),
    ...     ),
    ... ])
    """

    # ── Asignar colores automáticos a series sin color ─────────────────────
    default_colors = plt.cm.tab10.colors
    color_counter  = 0
    for s in series:
        if s.color is None:
            s.color = default_colors[color_counter % len(default_colors)]
        color_counter += 1

    fig, ax = plt.subplots(figsize=figsize, dpi=dpi)
    legend_handles: list[Line2D] = []

    # ── Iterar sobre las series ────────────────────────────────────────────
    for s in series:

        label = s.label if s.label is not None else _ashrae_label(s.fluid, s.composition)

        # Validar CycleDef si existe
        if s.cycle is not None:
            s.cycle.validate()

        # ── Curva de saturación ────────────────────────────────────────────
        p_min_sat = max(
            (_P_MIN_SAT.get(f, 0.0) for f in s.fluid), default=None
        ) or None

        try:
            s_liq, T_liq, s_vap, T_vap = _sat_curve(
                s.fluid, s.composition,
                n_points=s.n_points,
                p_min_sat=p_min_sat,
            )
        except Exception as exc:
            print(f"[AVISO] No se pudo calcular la curva de saturación para {s.fluid}: {exc}")
            continue

        ax.plot(s_liq, T_liq, color=s.color, linewidth=s.linewidth, linestyle=s.linestyle)
        ax.plot(s_vap, T_vap, color=s.color, linewidth=s.linewidth, linestyle=s.linestyle)

        # Cerrar la cúpula en el punto crítico
        if len(s_liq) and len(s_vap):
            try:
                T_crit_C = rprop(s.fluid, "Tcrit", s.composition, T=0, H=0)
                s_cp_liq = rprop(s.fluid, "S", s.composition, T=T_crit_C, Q=0)
                s_cp_vap = rprop(s.fluid, "S", s.composition, T=T_crit_C, Q=1)
                # El punto crítico está donde s_liq == s_vap == s_cp
                ax.plot([s_liq[-1], s_cp_liq], [T_liq[-1], T_crit_C],
                        color=s.color, linewidth=s.linewidth, linestyle=s.linestyle)
                ax.plot([s_vap[-1], s_cp_vap], [T_vap[-1], T_crit_C],
                        color=s.color, linewidth=s.linewidth, linestyle=s.linestyle)
            except Exception:
                pass

        # ── Ciclo ──────────────────────────────────────────────────────────
        if s.cycle is not None:
            cd = s.cycle

            # Ciclo teórico
            if cd._requires_theoretical():
                cycle_row = _compute_theoretical_cycle(
                    s.fluid, s.composition, cd.water_config
                )
                if cycle_row is not None:
                    _draw_theoretical_cycle(
                        ax, cycle_row,
                        color=s.color,
                        linewidth=cd.linewidth,
                        linestyle=cd.linestyle,
                        alpha=cd.alpha,
                    )

            # Ciclo experimental
            if cd._requires_experimental():
                try:
                    temperatures, entropies = _load_experimental_points(cd, label)
                except Exception as exc:
                    print(
                        f"[AVISO] No se pudieron cargar datos experimentales "
                        f"de '{label}': {exc}"
                    )
                    temperatures, entropies = {}, {}

                if temperatures and entropies:
                    # Cuando source=BOTH usar apariencia diferenciada
                    exp_ls    = cd.exp_linestyle if cd.source == CycleSource.BOTH else cd.linestyle
                    exp_alpha = cd.exp_alpha     if cd.source == CycleSource.BOTH else cd.alpha

                    _draw_experimental_cycle(
                        ax,
                        color=s.color,
                        temperatures=temperatures,
                        entropies=entropies,
                        marker=cd.marker,
                        markersize=cd.markersize,
                        linewidth=cd.linewidth,
                        linestyle=exp_ls,
                        alpha=exp_alpha,
                        series_label=label,
                    )

        # ── Entrada de leyenda ─────────────────────────────────────────────
        legend_handles.append(
            Line2D(
                [0], [0],
                color=s.color,
                linewidth=s.linewidth,
                linestyle=s.linestyle,
                label=label,
            )
        )

    # ── Límites, decoración ────────────────────────────────────────────────
    if s_min is not None or s_max is not None:
        ax.set_xlim(left=s_min, right=s_max)
    if T_min is not None or T_max is not None:
        ax.set_ylim(bottom=T_min, top=T_max)

    ax.set_xlabel(xlabel, fontsize=12)
    ax.set_ylabel(ylabel, fontsize=12)
    ax.set_title(title, fontsize=14, fontweight="bold")

    if grid:
        ax.grid(True, which="both", linestyle="--", linewidth=0.5, alpha=0.6)

    if legend_handles:
        ax.legend(
            handles=legend_handles,
            loc="upper left",
            framealpha=0.9,
            fontsize=10,
            title="Fluido",
            title_fontsize=10,
        )

    fig.tight_layout()
    return fig


# ---------------------------------------------------------------------------
# Ejemplo de uso directo
# ---------------------------------------------------------------------------

TEMPERATURE_COLS = {"T1": "AO", "T2": "AI", "T3": "BB", "T4": "AM"}
ENTROPY_COLS     = {"s1": "F", "s2": "G", "s3": "BH", "s4": "AS"}

if __name__ == "__main__":
    init_refprop()

    fig = plot_ts_diagrams(
        series=[
            FluidSeries(
                fluid       = ["PROPANE"],
                composition = [1.0],
                color       = "steelblue",
                cycle       = CycleDef(
                    source           = CycleSource.EXPERIMENTAL,
                    xlsx_path        = "res-exp.xlsx",
                    sheet_name       = "Hoja1",
                    row              = 4,
                    temperature_cols = TEMPERATURE_COLS,
                    entropy_cols     = ENTROPY_COLS,
                    linestyle        = "--",
                    alpha            = 0.7,
                ),
            ),

            FluidSeries(
                fluid       = ["PROPANE", "DME"],
                composition = [0.85, 0.15],
                cycle       = CycleDef(
                    source           = CycleSource.EXPERIMENTAL,
                    xlsx_path        = "res-exp.xlsx",
                    sheet_name       = "Hoja1",
                    row              = 10,
                    temperature_cols = TEMPERATURE_COLS,
                    entropy_cols     = ENTROPY_COLS,
                    linestyle        = "--",
                    alpha            = 0.7,
                ),
            ),

        ],
        s_min  = 0.5,
        s_max  = 2.5,
        T_min  = -50,
        T_max  = 120,
        title  = "Comparativa diagramas T-S de los ensayos experimentales",
    )

    fig.savefig("diagramas_TS/riley_exp_ts.png", bbox_inches="tight")
    print("Figura guardada")